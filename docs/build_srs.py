"""Build the Audio to PDF software requirements specification."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Audio_to_PDF_SRS_IEEE_29148.docx"


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    edges = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "D9D9D9")
        edges.append(edge)
    tc_pr.append(edges)


def cell_padding(cell, top=85, left=105, bottom=85, right=105):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("start", left), ("bottom", bottom), ("end", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = Inches(.72)
sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.82)
sec.right_margin = Inches(.82)
sec.header_distance = Inches(.30)
sec.footer_distance = Inches(.30)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.2)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.12
for name, size, before, after in (("Title", 22, 0, 13), ("Heading 1", 15, 15, 7),
                                  ("Heading 2", 12, 11, 5), ("Heading 3", 10.5, 8, 4)):
    style = styles[name]
    style.font.name = "Aptos Display" if name == "Title" else "Aptos"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
title_props = styles["Title"].element.get_or_add_pPr()
title_border = title_props.find(qn("w:pBdr"))
if title_border is not None:
    title_props.remove(title_border)

header = sec.header.paragraphs[0]
header.text = "Audio to PDF    |    Software Requirements Specification"
header.style = normal
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor(80, 80, 80)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("Working draft 0.1    |    Page ")
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
footer._p.append(fld)
for run in footer.runs:
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(80, 80, 80)


def para(text="", style=None):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p


def heading(text, level=1):
    return doc.add_heading(text, level=level)


def bullet(text):
    return para(text, "List Bullet")


def table(headers, rows, widths=None, compact=False):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    t.style = "Table Grid"
    if widths:
        for col, width in zip(t.columns, widths):
            col.width = Inches(width)
    for i, value in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = value
        shade(cell, "DFE6EF")
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    for r_index, row in enumerate(t.rows):
        for c_index, cell in enumerate(row.cells):
            if widths:
                cell.width = Inches(widths[c_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            borders(cell)
            cell_padding(cell, *(55, 75, 55, 75) if compact else (85, 105, 85, 105))
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(8 if compact else 8.5 if len(headers) > 3 else 9)
                    if r_index == 0:
                        run.bold = True
        if r_index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def reqs(section, rows):
    heading(section, 2)
    table(["ID", "Requirement", "Verification", "State"], rows, [.76, 4.45, 1.0, .52])


# Cover and document control
doc.add_paragraph("Audio to PDF", "Title")
p = para("Software Requirements Specification")
p.runs[0].font.size = Pt(16)
p.runs[0].font.bold = True
para("IEEE style structure based on ISO IEC IEEE 29148 2018")
para("Version 0.1    |    17 October 2026    |    Working draft")
para("Prepared for the Audio to PDF project team and student reviewers.")
para("This specification defines the intended behavior of a system that synchronizes spoken lectures with the exact source PDF page and region. It records both the current MVP and target behavior. It is a working project specification, not a certification of compliance with the IEEE standard.")

heading("Document control", 1)
table(["Item", "Value"], [
    ("Project", "Audio to PDF"),
    ("Document owner", "Project team"),
    ("Baseline inspected", "Local source in C:/Users/User/Desktop/Vscode/Audio-to-PDF on 17 October 2026"),
    ("Version", "0.1 working draft"),
    ("Status codes", "I = implemented in inspected code; P = planned or partially implemented; V = implemented but validation remains open"),
    ("Verification codes", "T = automated test; I = inspection; D = demonstration; M = measured evaluation"),
], [1.3, 5.43])

heading("Contents", 1)
for item in ("1 Introduction", "2 Overall description", "3 External interface requirements",
             "4 Functional requirements", "5 Data and information requirements",
             "6 Quality and constraint requirements", "7 Verification and acceptance",
             "8 Current baseline and open issues", "Appendix A Logical architecture",
             "Appendix B API inventory", "Appendix C References"):
    para(item)

doc.add_page_break()

heading("1 Introduction")
heading("1.1 Purpose", 2)
para("This SRS gives developers, reviewers, testers, and student users a shared basis for building and evaluating Audio to PDF. Each numbered requirement states an observable behavior and a verification method. Status describes the inspected MVP, while requirements describe the intended product.")
heading("1.2 Scope", 2)
para("A student uploads a lecture PDF and matching audio, or supplies an existing timestamped transcript. The system extracts the document structure, aligns spoken passages to source pages and blocks, and plays the original audio while showing the original PDF with timed underlines or thin orange boxes. Users can review and correct uncertain matches, revisit persistent analogy pages, and export a JSON timeline or compact MP4. Slide decks are supported when supplied as PDFs.")
para("The system does not promise correct automatic annotation for every PDF. Scanned documents, garbled mathematical symbols, unmatched narration, and low-confidence alignment must remain reviewable without a misleading mark. A complete generated video background is outside scope.")
heading("1.3 Audience and reading order", 2)
para("Product collaborators should read Sections 1, 2, and 8 first. Developers should continue with Sections 3–6 and the API appendix. Testers should use the numbered requirements and Section 7. Instructors or students evaluating lesson quality should focus on the review and acceptance criteria.")
heading("1.4 Terms and conventions", 2)
table(["Term", "Meaning"], [
    ("Source page", "A numbered page in the uploaded PDF, including a slide exported to PDF."),
    ("Document block", "A positioned heading, paragraph, definition, theorem, example, exercise, equation, or figure on a source page."),
    ("Grounding", "Selecting the source page and block supported by a spoken passage."),
    ("Point", "A short, timestamped transcript passage with its page, block, action, evidence, and review state."),
    ("Annotation", "A timed underline or thin orange box drawn above a source page."),
    ("Whiteboard", "A virtual lesson page for a spoken analogy; the uploaded PDF is unchanged."),
    ("Abstain", "Keep a page visible and request review without adding a new annotation."),
], [1.27, 5.46])
heading("1.5 Requirements language", 2)
para("“Shall” denotes a requirement. Normalized regions use [x, y, width, height] values between zero and one relative to the source page. Times are seconds from the beginning of the narration. Requirement IDs are stable references for tests and future revisions.")

heading("2 Overall description")
heading("2.1 Product perspective", 2)
para("Audio to PDF is a local web application with a browser player and a Python FastAPI server. It stores each lesson under a local data directory. The source PDF is rendered into page images; annotations are separate timed events, preserving the PDF. The same lesson timeline drives playback and MP4 export.")
heading("2.2 User classes", 2)
table(["User", "Primary tasks"], [
    ("Student", "Upload or open a lesson, listen and seek, follow marks, browse source and analogy pages, download a video."),
    ("Teacher or lesson editor", "Review uncertain alignments, correct pages and regions, edit equation cues and whiteboards, approve explanatory visuals."),
    ("Local operator", "Install dependencies, start the server, configure optional Gemini access, inspect processing and export failures."),
], [1.42, 5.31])
heading("2.3 Operating environment and dependencies", 2)
para("The inspected implementation targets Python 3.11 or later on Windows and a modern browser. It uses FastAPI for the API, pdfplumber and pypdfium2 for PDF extraction and rendering, faster-whisper for local transcription, RapidOCR for fallback OCR, FastEmbed for local semantic retrieval, and FFmpeg for MP4 encoding. Gemini support is optional and server-side. Fresh speech and embedding model setups may require one-time downloads.")
heading("2.4 Assumptions and constraints", 2)
bullet("Uploads consist of one PDF and one related audio file; transcript JSON is optional. Audio and PDF may cover different lengths, and some speech may have no exact source region.")
bullet("A source page stays stationary during explanation. The player changes pages directly when a grounded event calls for it; automatic zoom and continuous scrolling are excluded.")
bullet("PDF annotations are restricted to underlines and orange boxes. Analogy sketches are stored on separate whiteboard pages.")
bullet("A free-tier Gemini key may have provider quotas. Equation grounding and the current core alignment must work locally without Gemini calls.")
bullet("The current application is local and has no user accounts, collaboration service, or remote deployment requirement.")

heading("3 External interface requirements")
heading("3.1 User interface", 2)
reqs("3.1.1 Browser interface", [
    ("UI-01", "The interface shall present a lesson list and a New lesson form with PDF, audio, and optional transcript inputs.", "D", "I"),
    ("UI-02", "The player shall show the original page, a separate annotation layer, current topic and review state, audio controls, and a timed point list.", "D", "I"),
    ("UI-03", "The player shall provide previous and next lesson-page controls, a page selector, and a Resume audio page control after manual browsing.", "D", "I"),
    ("UI-04", "The editor shall allow a user to correct a point's source page, document block, action, concept, and target region.", "D/T", "I"),
    ("UI-05", "The interface shall distinguish uncertain points from verified annotations and show processing or export errors in readable language.", "D/T", "V"),
    ("UI-06", "The interface shall provide JSON and compact MP4 export controls, progress, file size, and a download action.", "D", "I"),
])
heading("3.2 File and software interfaces", 2)
reqs("3.2.1 Input and output files", [
    ("IF-01", "The system shall accept PDF input with a valid PDF signature and audio in MP3, M4A, WAV, MP4, OGG, or FLAC format, up to 120 MB per file in the local MVP.", "T", "I"),
    ("IF-02", "The system shall accept optional UTF-8 JSON containing a segments array and shall reject invalid transcript structure.", "T", "I"),
    ("IF-03", "The system shall store the source PDF, original audio, transcript, rendered pages, lesson events, and any whiteboard assets per lesson.", "I/T", "I"),
    ("IF-04", "The system shall expose lesson creation, retrieval, editing, rebuild, and export through local HTTP endpoints listed in Appendix B.", "T", "I"),
    ("IF-05", "The system shall keep Gemini credentials on the server and shall not return them in lesson JSON or browser assets.", "I/T", "I"),
])
heading("3.3 Hardware and communications", 2)
para("No special hardware is required beyond a computer able to render PDFs, run local speech and embedding models, and encode H.264. Browser and server communicate over localhost HTTP. The system makes external requests only for one-time model acquisition or optional Gemini operations when configured.")

heading("4 Functional requirements")
reqs("4.1 Lesson ingestion and source extraction", [
    ("FR-01", "The system shall create a distinct lesson identifier and process a validated upload asynchronously.", "T", "I"),
    ("FR-02", "The system shall render every source PDF page at a consistent size and retain its original page number and dimensions.", "T/I", "I"),
    ("FR-03", "The system shall extract positioned words and visible equations or figures from searchable PDF pages.", "T/I", "I"),
    ("FR-04", "When embedded text is insufficient, the system shall attempt local OCR and record that OCR supplied the text.", "T", "V"),
    ("FR-05", "The system shall remove repeated headers and footers from grounding candidates without altering the rendered page.", "T", "V"),
    ("FR-06", "The system shall create stable, positioned blocks for headings, definitions, theorems, examples, exercises, prose, equations, and figures.", "T", "V"),
    ("FR-07", "The system shall classify unreadable mathematical text conservatively and send it for review instead of inferring symbols that were not extracted.", "T", "V"),
])
reqs("4.2 Audio analysis and alignment", [
    ("FR-08", "The system shall preserve the original narration and derive a transcript with word start and end times when no transcript is supplied.", "T/D", "I"),
    ("FR-09", "The system shall group transcript segments into short, timed passages using sentence and topic cues rather than fixed 15-second windows.", "T", "I"),
    ("FR-10", "The system shall identify explicit section, example, exercise, and heading references in the narration.", "T", "V"),
    ("FR-11", "The system shall retrieve candidate source blocks using exact terms, section numbers, lexical evidence, and local semantic similarity when the model is available.", "T", "V"),
    ("FR-12", "The system shall align passages across the full lecture sequence while permitting an explicit return to an earlier page.", "T", "V"),
    ("FR-13", "Each aligned point shall store its selected page, block identifier where supported, evidence, confidence, start and end times, action, and review state.", "T/I", "I"),
    ("FR-14", "When no source block is supported strongly enough, the system shall retain the last verified page, add no new mark, and put the point in review.", "T", "V"),
    ("FR-15", "The core document alignment shall function without a Gemini API key and shall make zero Gemini calls.", "T/I", "I"),
])
reqs("4.3 Annotation and equation focus", [
    ("FR-16", "The system shall underline verified text within its selected document block using a thin line beneath the spoken phrase.", "D/T", "V"),
    ("FR-17", "The system shall draw a thin orange rectangular border around a verified equation or figure and shall avoid obscuring its contents.", "D/T", "V"),
    ("FR-18", "The system shall withhold automatic boxes around garbled equations or poorly located figures and request review.", "T", "V"),
    ("FR-19", "For searchable equations, a local term mapper shall be able to identify and time equation-part cues such as new weight, old weight, learning rate, and derivative.", "T", "I"),
    ("FR-20", "At most one active equation focus box shall move to the currently spoken term; ungrounded terms shall hold the whole equation or enter review.", "D/T", "V"),
    ("FR-21", "Equation focus events shall be editable by cue time, equation, term, and region; rebuild shall preserve deliberate corrections.", "T/D", "I"),
])
reqs("4.4 Analogy interpretation and whiteboards", [
    ("FR-22", "The system shall distinguish spoken analogies from explanations of source figures and shall treat uncertain classifications as ordinary narration requiring review.", "T", "P"),
    ("FR-23", "A drawable analogy shall use a persistent virtual whiteboard inserted after its related source page, without changing the PDF.", "D/T", "I"),
    ("FR-24", "The whiteboard shall reveal safe drawing steps in audio order and shall remain available after the narration moves on.", "D/T", "I"),
    ("FR-25", "The drawing schema shall permit only line, curve, arrow, point, rectangle, ellipse, and text with valid normalized coordinates and bounded labels.", "T", "I"),
    ("FR-26", "The system may use one cached, text-free generated illustration for an analogy that cannot be explained by simple shapes, with timed editable labels and strokes above it.", "D/T", "P"),
    ("FR-27", "Optional Gemini planning and image generation shall run server-side, share a persistent limit of eight calls per lesson, and stop further calls on quota exhaustion.", "T/I", "I"),
    ("FR-28", "Without Gemini access, the system shall offer local hill, flow, input-output, network, and momentum templates and route other analogies to review.", "T/D", "P"),
    ("FR-29", "For newly created lessons, automatic analogy page insertion shall remain paused until the source-grounding acceptance criteria in Section 7 pass.", "I", "I"),
])
reqs("4.5 Playback, editing, and persistence", [
    ("FR-30", "The player shall display a full source page without automatic zoom or continuous scroll and shall change pages directly at grounded cues.", "D", "I"),
    ("FR-31", "Saved PDF annotations shall reappear on return visits to their page; whiteboard drawings shall remain in lesson-page order.", "D/T", "I"),
    ("FR-32", "Seeking or changing playback speed shall update the visible page and active annotation from the saved timeline.", "D/T", "V"),
    ("FR-33", "Manual page browsing shall pause playback; Resume audio page shall restore synchronization.", "D/T", "I"),
    ("FR-34", "A reviewer shall be able to save a corrected page, block, region, action, or concept and mark a point reviewed.", "D/T", "I"),
    ("FR-35", "A reviewer shall be able to edit and approve whiteboard steps and remove an incorrect generated illustration.", "D/T", "I"),
    ("FR-36", "Rebuilding alignment shall retain manual corrections and existing whiteboards, assign unique point IDs, and validate block references against the new map.", "T", "V"),
    ("FR-37", "Existing lessons without virtual-page or equation-focus fields shall open without destructive migration.", "T", "V"),
])
reqs("4.6 Export and jobs", [
    ("FR-38", "The system shall export the editable lesson timeline as JSON with source pages, points, whiteboards, and equation focus events.", "T", "I"),
    ("FR-39", "The system shall export the complete narration and synchronized page, annotation, and whiteboard timeline as a 1280 by 720 MP4.", "T/D", "I"),
    ("FR-40", "MP4 encoding shall use H.264 at 24 frames per second with CRF 28, yuv420p, mono AAC at 64 kbps, and no zoom or transition effects.", "I/T", "I"),
    ("FR-41", "Video frames shall stream to FFmpeg without storing thousands of intermediate images; a bundled provider may supply FFmpeg on Windows.", "I/M", "I"),
    ("FR-42", "An export job shall expose queued, processing, ready, or failed state with progress and a final download link or readable error.", "T/D", "I"),
    ("FR-43", "A repeat export of an unchanged lesson shall use the cached file; changing an annotation or whiteboard shall change the export hash.", "T", "I"),
])

heading("5 Data and information requirements")
heading("5.1 Logical entities", 2)
table(["Entity", "Required fields and relationships"], [
    ("Lesson", "ID, title, audio filename, duration, grounding version, source pages, lesson pages, points, whiteboards, equation focus, statistics."),
    ("Source page", "PDF page number, image path, width, height, positioned words, text quality, document blocks, visual regions."),
    ("Document block", "Stable ID, page, kind, text, normalized region, line regions, source quality, section and parent references."),
    ("Timed point", "Unique ID, start/end, transcript text and words, page/page ID, block ID when supported, action, region, evidence, confidence, review and edited flags."),
    ("Whiteboard", "Stable ID, placement, title, start/end, template or illustration, safe commands, review status."),
    ("Equation focus", "Unique event ID, time window, page, equation and term IDs, region, evidence, confidence, review and edited flags."),
    ("Export", "Lesson timeline hash, job state, progress, file size, cached MP4 path."),
], [1.28, 5.45])
heading("5.2 Data integrity and retention", 2)
reqs("5.2.1 Integrity requirements", [
    ("DR-01", "All point IDs within one lesson shall be unique, including after alignment rebuild and whiteboard preservation.", "T", "V"),
    ("DR-02", "Every saved source-region coordinate shall be finite and lie within the normalized page bounds.", "T", "V"),
    ("DR-03", "Every saved block ID shall refer to a block on the point's selected source page, or be null when no reliable block exists.", "T", "V"),
    ("DR-04", "Transcript and event times shall be nonnegative, ordered, and bounded by the audio duration with a documented tolerance.", "T", "V"),
    ("DR-05", "Lesson edits shall be written atomically so interruption does not leave a partially written lesson JSON.", "T/I", "I"),
    ("DR-06", "The product shall keep uploaded material and derived lessons on the local machine unless the user deliberately exports or shares a file.", "I", "I"),
])

heading("6 Quality and constraint requirements")
reqs("6.1 Accuracy and safety", [
    ("QR-01", "On the two supplied lectures, automatic page assignment shall reach at least 90 percent accuracy over at least 60 human-labeled timed passages.", "M", "V"),
    ("QR-02", "Among automatically marked passages with clear source counterparts, block selection shall reach at least 80 percent precision and at least 80 percent coverage.", "M", "P"),
    ("QR-03", "Reported evaluation shall separate incorrect marks, correct marks, and abstentions so conservative withholding cannot conceal low coverage.", "M", "P"),
    ("QR-04", "The calculus sample shall move to page 2 for Divergence Theorem, page 3 for continuation and exercises, and page 4 for later exercises at human-checked times.", "M/D", "V"),
    ("QR-05", "Uncertain OCR, mathematical notation, or unsupported narration shall produce no automatic visual mark and shall remain visible in review.", "T", "V"),
])
reqs("6.2 Performance, portability, and usability", [
    ("QR-06", "The system shall reuse downloaded local models and cached exports, and shall report processing progress or failure for long jobs.", "T/D", "V"),
    ("QR-07", "The complete 19-minute sample MP4 shall match narration duration within one second and target less than 60 MB while retaining legible text at 720p.", "M/D", "V"),
    ("QR-08", "The exported MP4 shall play in Chrome, Edge, and Windows Media Player with narration synchronized to page and drawing events.", "D", "V"),
    ("QR-09", "The interface shall keep annotation coordinates aligned when the viewer resizes and shall expose controls through native labels and buttons.", "D/T", "V"),
    ("QR-10", "A page with weak or no match shall keep the most recent verified page visible rather than jump on a speculative semantic similarity.", "T/D", "V"),
])
reqs("6.3 Security and operational constraints", [
    ("QR-11", "The server shall validate lesson IDs, file types, upload size, region bounds, drawing commands, and export identifiers before use.", "T/I", "V"),
    ("QR-12", "The server shall not send Gemini API keys to the browser or include them in exported lesson data.", "T/I", "I"),
    ("QR-13", "Gemini quota errors shall leave locally created lessons usable and mark skipped analogy work for review.", "T", "V"),
    ("QR-14", "The core source-grounding path shall make no external AI request and shall remain functional when Gemini is unset.", "T/I", "I"),
])

heading("7 Verification and acceptance")
heading("7.1 Evaluation corpus and method", 2)
para("Use the supplied optimization lecture and the supplied MA2024 calculus lecture, each with its matching audio. Label at least 60 timed passages across page transitions, repeated vocabulary, equations, figures, explicit backward references, and narration absent from the source. Separately label source blocks for passages with clear PDF counterparts. Measure page accuracy, automatic-block precision, grounded coverage, incorrect-mark rate, and abstention rate. Keep labels and scoring scripts under version control.")
heading("7.2 Required scenario tests", 2)
for text in (
    "An uploaded searchable slide deck maps a spoken formula to the correct page and a specific visible region; unreadable formula symbols create review without a guessed term.",
    "A scanned page uses local OCR; malformed or empty OCR output does not crash lesson creation and yields a reviewable hold.",
    "An explicit spoken return to an earlier slide restores its saved annotations, then follows the next grounded event.",
    "Seeking, speed change, manual page browsing, correction, and lesson reopen retain the same timeline interpretation.",
    "A rebuild preserves edited points and whiteboards, keeps unique IDs, and does not carry stale block IDs onto different text.",
    "Browser playback and MP4 output agree at sampled times for page, underline, orange box, whiteboard stage, and equation focus.",
    "The complete sample export reports render time and file size, matches audio duration within one second, and a repeated request returns the cached file.",
): bullet(text)
heading("7.3 Acceptance gate", 2)
para("Do not describe arbitrary uploads as accurately grounded until QR-01 through QR-05 pass on a labeled evaluation set. Failure cases must remain easy to correct, and a reviewed lesson should be checked by a person before sharing as instructional material.")

heading("8 Current baseline and open issues")
heading("8.1 Inspected implementation", 2)
para("The current code contains the local structured PDF map, OCR fallback, sentence-based transcript grouping, full-lecture page and block alignment, conservative annotation gating, browser editor, alignment rebuild endpoint, persistent legacy whiteboards, optional Gemini analogy module, local equation-focus module, and JSON/MP4 export. Fresh lessons currently create source PDF pages and no new whiteboards or equation-focus events. Existing saved lessons may still contain those features.")
heading("8.2 Known validation gaps at this document version", 2)
para("The latest local run passed all 13 tests. On 64 timed page landmarks from the calculus and optimization lectures, 60 were correct (93.75 percent); two landmarks fell in transcript gaps, and two optimization slide boundaries were wrong. Among 19 labeled block passages, 16 received automatic marks and 15 selected the labeled block: 84.2 percent coverage and 93.75 percent precision on this small set. These figures describe the supplied samples, not arbitrary uploads.")
para("The page-three calculus title now resolves to “Exercises for Lecture 7.” Rebuild regression tests now check unique IDs, preserved corrections and whiteboards, and stale block references. The block evaluation still has fewer than 60 labels; scanned-PDF/OCR failure handling, manual browser flows, and browser-to-MP4 parity need broader validation. These remain open acceptance items rather than completed claims.")
heading("8.3 Priority order", 2)
table(["Priority", "Work item", "Exit condition"], [
    ("1", "Resolve remaining optimization slide-boundary errors and strengthen the page labels.", "QR-01 stays above 90 percent on a larger set."),
    ("2", "Test alignment rebuild through the API with saved lessons and multiple corrections.", "FR-36 and DR-01 to DR-03 pass end to end."),
    ("3", "Expand block, OCR, unsupported-speech, and parity tests.", "QR-02, QR-03, QR-05 pass; scenarios in 7.2 pass."),
    ("4", "Resume automatic analogy creation after source grounding is dependable.", "FR-22 through FR-29 demonstrated without misleading source marks."),
], [.76, 3.95, 2.02])

heading("Appendix A Logical architecture")
table(["Stage", "Input", "Output"], [
    ("1 Source ingestion", "PDF and audio or transcript", "Validated local lesson directory and job"),
    ("2 Document map", "PDF text, glyph geometry, page image, OCR if needed", "Typed blocks, regions, page titles, source quality"),
    ("3 Audio analysis", "Original audio or supplied transcript", "Word times and sentence-level points"),
    ("4 Grounding", "Timed points and full document map", "Page, block, evidence, confidence, action or review"),
    ("5 Optional visuals", "Grounded equations; classified analogies", "Equation-focus events; persistent whiteboard commands or one cached illustration"),
    ("6 Lesson player", "Original audio and saved events", "Stationary source page, timed overlay, navigation and correction"),
    ("7 Export", "Same saved timeline", "JSON or 720p H.264/AAC MP4"),
], [1.25, 2.7, 2.78])
para("The transcript analyzer describes what is being said and when. The grounding engine determines which visible source block supports it. The renderer executes saved events; it does not decide new matches during playback. Any future analogy storyboard uses local word times as the timing authority. The PDF page and its annotations stay intact when an inserted whiteboard is visited.")

heading("Appendix B API inventory")
table(["Method and route", "Purpose"], [
    ("GET /api/health", "Local server health."),
    ("GET /api/lessons", "List saved lessons."),
    ("POST /api/lessons", "Upload PDF, audio, and optional transcript; queue processing."),
    ("GET /api/jobs/{lesson_id}", "Read creation or alignment-rebuild status."),
    ("GET /api/lessons/{lesson_id}", "Read lesson metadata, blocks, events, and review data."),
    ("POST /api/lessons/{lesson_id}/alignment/rebuild", "Re-run local PDF alignment."),
    ("GET /api/lessons/{lesson_id}/audio", "Stream original audio."),
    ("GET /api/lessons/{lesson_id}/pages/{number}", "Get rendered source page."),
    ("PATCH /api/lessons/{lesson_id}/points/{point_id}", "Save a point correction."),
    ("PATCH /api/lessons/{lesson_id}/whiteboards/{board_id}", "Edit or approve a whiteboard."),
    ("POST /api/lessons/{lesson_id}/equation-focus/rebuild", "Rebuild local equation focus."),
    ("PATCH /api/lessons/{lesson_id}/equation-focus/{event_id}", "Correct an equation cue."),
    ("GET /api/lessons/{lesson_id}/export.json", "Download event timeline."),
    ("POST /api/lessons/{lesson_id}/exports", "Queue or reuse an MP4."),
    ("GET /api/exports/{job_id}", "Read MP4 progress or failure."),
    ("GET /api/lessons/{lesson_id}/exports/{export_id}.mp4", "Download current cached MP4."),
], [3.55, 3.18], compact=True)

heading("Appendix C References")
para("[1] ISO/IEC/IEEE 29148:2018, Systems and software engineering — Life cycle processes — Requirements engineering, ISO catalog: https://www.iso.org/standard/72089.html. This document uses an IEEE-style SRS structure and testable requirement statements; it does not reproduce the standard.")
para("[2] Audio to PDF project README and inspected source modules: backend/analysis.py, backend/grounding.py, backend/analogies.py, backend/equation_focus.py, backend/video_export.py, backend/app.py, and frontend/. Local repository, inspected 17 October 2026.")
para("[3] Local development tests: tests/test_grounding.py and tests/test_equation_focus.py. The status in Section 8 records the latest run available when this document was drafted.")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
