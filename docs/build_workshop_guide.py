"""Build the sourced AI Builder Cup workshop guide."""

from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "AI_Builder_Cup_Workshop_Guide.docx"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = Inches(.78)
sec.bottom_margin = Inches(.72)
sec.left_margin = Inches(.82)
sec.right_margin = Inches(.82)
sec.header_distance = Inches(.35)
sec.footer_distance = Inches(.35)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Aptos'
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(25, 31, 39)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.13
for name, size, before, after in [('Title', 19, 0, 11), ('Heading 1', 14, 16, 7), ('Heading 2', 11.5, 12, 5), ('Heading 3', 10.5, 8, 3)]:
    st = styles[name]
    st.font.name = 'Aptos Display' if name == 'Title' else 'Aptos'
    st.font.size = Pt(size)
    st.font.bold = name != 'Title'
    st.font.color.rgb = RGBColor(0,0,0)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True
if 'Source Note' not in styles:
    s = styles.add_style('Source Note', WD_STYLE_TYPE.PARAGRAPH)
else: s = styles['Source Note']
s.font.name='Aptos'; s.font.size=Pt(8.5); s.font.color.rgb=RGBColor(78,88,101)
s.paragraph_format.space_after=Pt(5)

def p(text='', bold_lead=None, style=None):
    x=doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        x.add_run(bold_lead).bold=True
        x.add_run(text[len(bold_lead):])
    else: x.add_run(text)
    return x

def b(text):
    x=doc.add_paragraph(style='List Bullet')
    x.add_run(text)
    x.paragraph_format.space_after=Pt(4)
    return x

def h(text, level=1): doc.add_heading(text, level=level)

def src(text): p(text, style='Source Note')

def cell_shade(cell, color):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),color); tcPr.append(shd)

def cell_border(cell):
    tcPr=cell._tc.get_or_add_tcPr(); borders=tcPr.first_child_found_in('w:tcBorders')
    if borders is None: borders=OxmlElement('w:tcBorders'); tcPr.append(borders)
    for side in ('top','left','bottom','right'):
        x=OxmlElement('w:'+side); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'D9D9D9'); borders.append(x)

def table(headers, rows, widths=None, font=8.8):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for j,v in enumerate(headers):
        c=t.rows[0].cells[j]; c.text=v; cell_shade(c,'203B63')
        for run in c.paragraphs[0].runs: run.bold=True; run.font.color.rgb=RGBColor(255,255,255)
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        for j,v in enumerate(row):
            cells[j].text=str(v)
            if i%2: cell_shade(cells[j],'F3F6FA')
    for row in t.rows:
        for j,c in enumerate(row.cells):
            if widths: c.width=Inches(widths[j])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_border(c)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(1)
                para.paragraph_format.line_spacing=1.03
                for run in para.runs: run.font.name='Aptos'; run.font.size=Pt(font)
            tc=c._tc; tcPr=tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                x=OxmlElement('w:'+side); x.set(qn('w:w'),'80'); x.set(qn('w:type'),'dxa'); mar.append(x)
            tcPr.append(mar)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

title=doc.add_paragraph(style='Title'); title.add_run('AI Builder Cup Workshop Guide for Audio to PDF')
p('Workshop analysis, grading playbook, submission requirements and project decisions')
src('Prepared 18 October 2026  |  Source: 51-page Untitled document (6).pdf containing two workshop transcripts; official organizer pages checked 18 October 2026.')
p('This guide captures the substantive advice from both workshops and applies it to the Audio to PDF prototype. It distinguishes published requirements, workshop guidance and project recommendations. It does not predict a score or guarantee selection.')

h('1  Read this first')
p('Eligibility is the immediate decision. You confirmed that you are currently enrolled as a student. Both workshops say this edition is not open to students, and the current organizer page states that a team containing a student may be disqualified even after shortlisting. Registration or dashboard access should not be treated as an eligibility ruling. Ask Hack2skill for a written determination using the contact in section 9 before relying on this competition as the project’s submission route.')
src('Workshop 1, p. 4, 05:57–06:25; Workshop 2, p. 37, timestamp not supplied; official overview, https://aibuildercup.com/ .')
p('If the organizer confirms eligibility, the fastest scoring route is to show a real, reachable product that accurately connects engineering audio to the correct PDF page and block, uses Google AI for a meaningful part of that task, and makes uncertainty visible instead of drawing a misleading mark. The working deployment, evidence and three-minute demo matter as much as the architecture explanation.')

h('2  Highest priority actions')
table(['Priority','Action for Audio to PDF','Why it matters'],[
('Immediate','Obtain a written eligibility answer; check the live dashboard for exact cutoff and team status.','A technically strong submission can be disqualified before scoring.'),
('Before submission','Put the full evaluator-facing app on Cloud Run or Firebase and test it from a clean browser.','The published rules require a working Google Cloud deployment.'),
('Before submission','Make Google AI solve a visible, material grounding problem; document input, output, validation and fallback.','Gen AI contribution is central to the 40% technical criterion.'),
('Before submission','Demonstrate correct source-page transitions, exact block grounding and abstention on uncertain passages using two real lectures.','This is the product’s central user value and its strongest technical evidence.'),
('Before submission','Prepare the proposal PDF, public demo video, accessible repository and evaluator guide; verify every link.','Judges need to understand and operate the submitted system.'),
('After core works','Show measured latency, processing time, accuracy, file size and cost; explain known limits.','Evidence supports feasibility and differentiates the design.'),
], [1.05,3.45,2.36], 8.7)
src('Rubric and mandatory submission items: https://aibuildercup.com/themes.html ; Workshop 1 pp. 3–4, 10–11, 20–22; Workshop 2 pp. 43–46. The ordering is this guide’s recommendation.')

h('3  Competition requirements and dates')
p('The current organizer overview lists 2–4 team members, age 21 or over, the JAPAC region and a prototype submission deadline of 18 October 2026. It lists evaluation from 19 October to 6 November. The page gives no submission cutoff time or timezone; verify these in the participant dashboard. Team formation is shown as 14 October on the current site, while the introductory workshop says 11 October and the later Q&A says it was reopened to 16 October. Treat the dashboard as the operational source of truth.')
src('Workshop 1, pp. 2 and 35, approximately 03:00 and 52:30–52:55; Workshop 2, pp. 36–37, timestamp not supplied; official overview, https://aibuildercup.com/ .')
p('The official challenge page requires a proposal deck converted to PDF, a functional prototype using Google AI technology deployed on Google Cloud, and a publicly accessible three-minute demo on YouTube, Vimeo or Google Drive. Workshop speakers also request an accessible GitHub repository and English-language code, documentation and presentation materials. Keep the deployed prototype working throughout evaluation; the first workshop warns that unavailable APIs or integrations can prevent evaluation.')
src('Workshop 1, p. 3, 04:32–05:31; Workshop 2, pp. 43–45, timestamp not supplied; official challenges, https://aibuildercup.com/themes.html .')
p('A workshop speaker says a commercial product may be entered and describes a small public evidence repository if proprietary code cannot be shared. This is Q&A guidance, not a blanket exemption from the submission form. The public challenge page asks for a publicly accessible repository; check the dashboard before relying on a private repo or partial code release.')
src('Workshop 1, pp. 6–7, 08:41–10:08; Workshop 2, p. 45, timestamp not supplied.')

h('4  Theme choice and problem statement')
table(['Workshop theme','Main examples from the introduction','Fit for this product'],[
('Banking finance insurance','Risk, fraud, secure service, process automation and personalization.','Low'),
('Retail and commerce','Shopping, returns, support, inventory, demand planning and personalization.','Low'),
('Manufacturing','Quality, defects, supply chains, efficiency and reliability.','Indirect'),
('Media content digital experiences','Content creation, personalization, immersive education and multilingual access.','Strongest'),
('Future of work productivity','Knowledge access, collaboration, decision support and workflow improvements.','Possible secondary framing'),
('Sustainability social impact','Accessibility, community needs, climate and resource efficiency.','Possible benefit, weaker primary fit'),
],[1.75,3.45,1.66],8.5)
src('Workshop 2, pp. 38–42, timestamp not supplied; official themes, https://aibuildercup.com/themes.html .')
p('Recommended problem statement: Engineering students listening to long, clear audio explanations lose track of the precise equation, figure or passage being discussed in a large lecture PDF. Audio to PDF converts narration and the original source into a synchronized, reviewable visual lesson without generating a new video background for every frame. The strongest claim is improved source visibility; a learning-gain claim requires student testing.')
p('Use Media, Content & Digital Experiences as the primary category because the workshop explicitly discusses immersive education and personalized teaching there. State one learner, one source format and one task in the deck. The workshop emphasizes that the six themes are broad and expects teams to define a specific problem within a chosen theme.')
src('Workshop 2, pp. 40–44, timestamp not supplied; official theme wording, https://aibuildercup.com/themes.html .')

h('5  Grading playbook')
p('The published weights below are official. The project evidence and actions are this guide’s assessment of the checked README and code, not a judge’s score.')
table(['Criterion','What judges assess','Current evidence','Action and proof to add'],[
('Technical merit and Gen AI 40%','Meaningful Gen AI, sound implementation, scaling and sustainability.','Local word timing, document blocks, OCR fallback, full-lecture alignment, review flow, equation focus and MP4 exist. Gemini is optional for analogies; core alignment makes no Gemini calls.','Add a bounded, inspectable Google AI step to the actual grounding workflow, such as disambiguating hard candidate blocks using transcript context and page evidence. Validate its outputs locally, compare with a local baseline and show error handling. Deploy the working path on Google Cloud.'),
('Problem alignment and impact 25%','Clear theme fit and meaningful value to intended users.','The product targets engineering learners and large lecture PDFs. Two real lectures are available.','Show a 30–60 second learner task, a correctly located equation and a passage with no reliable source match. Interview or test learners; report observed improvement without inventing results.'),
('Innovation and creativity 25%','Differentiation and a Gen AI use beyond a superficial prompt.','Source PDF remains the background; editable timed events and conservative grounding distinguish the experience from a generic generated overview.','Show the event representation, precise source location, persistent annotations, correction and compact export. Compare fairly with plain audio, static PDF and summary video; avoid claiming no prior product exists.'),
('User experience and design 10%','Usability, accessibility and coherent integration.','Player has stationary PDF pages, underlines, boxes, review, seeking, page list and MP4 export.','Run a first-time-user walkthrough, verify responsive layout, readable equations, audible narration, obvious review state and low-friction evaluator access.'),
],[1.23,1.48,1.75,2.40],8.3)
src('Official rubric: https://aibuildercup.com/themes.html . Current project: README.md and backend analysis, grounding, analogy and export modules, checked 18 October 2026.')

h('6  What the workshops recommend technically')
h('Build the smallest coherent system',2)
p('The introductory speaker recommends starting with the problem and a simple system consisting of an AI component, runtime and data store. Add caches, load balancing, specialist databases or an agent framework when the use case needs them. Gemini SDK or ADK is a choice, not itself a scoring criterion. RAG can be built or managed; it is useful when retrieval needs justify it, not a mandatory architecture label.')
src('Workshop 2, pp. 47–50, timestamp not supplied; Workshop 1, pp. 9–10, 13:02–15:01.')
h('Ground outputs and show evidence',2)
p('The Q&A speaker says judges may inspect the repository, data source and guardrails to see whether outputs are grounded. Provide representative inputs, expected behavior and any test credentials. For this product, each automatic mark should be traceable to transcript words, a PDF block, its page and confidence. When there is no credible match, hold the page and route that point to review. This is stronger evidence than a plausible-looking underline with no source justification.')
src('Workshop 1, pp. 10–11, 15:48–17:42. Product recommendation is an inference from that guidance.')
h('Use Google AI substantively and safely',2)
p('The official criteria allow Google models such as Gemini or Gemma, or specified Google agentic platforms, with deployment on Cloud Run or Firebase. The workshop Q&A emphasizes a meaningful Gemini role even in a hybrid system that also uses local models. It does not say every operation must be generative. A suitable product design is local timestamps and geometry for precision, with a constrained Gemini pass on ambiguous topic, figure or block candidates, followed by deterministic validation and human review. Run API calls server-side, cap use, cache results and handle quota failure.')
src('Workshop 1, pp. 9–10 and 13–14, 13:02–15:01 and 20:12–22:07; Workshop 2, pp. 43 and 45, timestamp not supplied; official requirements, https://aibuildercup.com/themes.html . The proposed Gemini task is this guide’s recommendation, not a workshop mandate.')
h('Keep the evaluator path reliable',2)
p('The Q&A allows a Cloud Run or Firebase URL as the visible submitted entry point, even if some back-end components are elsewhere. That arrangement adds failure points and should be tested end to end. The speaker says Cloud Run may scale to zero and suggests measuring actual slow requests before tuning timeouts; do not buy idle capacity merely for a prototype. State known quotas or interaction limits so judges do not unintentionally exhaust the demo. Keep secrets out of the repository, video, deck and deployment commands; the Q&A calls exposed credentials a disqualification risk.')
src('Workshop 1, pp. 7–8, 15–19 and 29–31, 10:16–12:40, 22:47–29:30 and 43:44–47:08.')
h('Choose data and models for the workload',2)
p('The Q&A accepts hybrid local and Google AI processing and synthetic or public demo data when clearly labeled. It advises choosing Firestore, SQL or other storage according to data shape, scale and latency rather than adding services for appearance. For Audio to PDF, the immediate workload is per-lesson PDF geometry, audio, event timelines and export files. A distributed vector database is not needed solely to impress judges; explain a future scale path after the core works.')
src('Workshop 1, pp. 13–14, 24 and 31–34, 20:12–22:07, 36:07–36:23 and 47:13–51:28; Workshop 2, pp. 47–50, timestamp not supplied. Workload conclusion is project analysis.')

h('7  Evidence to collect for this project')
table(['Question a judge may ask','Evidence to prepare','Current caution'],[
('Does it locate the right source?','Hand-labeled passages from both lectures; page accuracy, block accuracy, coverage and incorrect-mark rate reported separately. Show transition and return cases.','README reports 60/64 labeled page landmarks and 15/16 correctly selected marked blocks on a small sample; two optimization boundaries remain wrong.'),
('Does Google AI improve the result?','Run the same hard cases with and without the proposed Gemini grounding step; show a few accepted and rejected model outputs.','Optional analogy generation alone does not demonstrate Google AI improving core PDF alignment.'),
('Is the equation focus precise?','Show a short term-by-term explanation and cue timing; include mismatches and scanned-equation behavior.','The 21/21 labeled local term-cue result is sample-specific and should not be presented as general accuracy.'),
('Can someone actually use it?','Clean-browser upload and play test, seeking, visible page title, review/edit action, and sample evaluator guide.','Large uploads and first local model download may be slow.'),
('Is it efficient and feasible?','Processing time per lecture, Cloud cost estimate after deployment, MP4 time/size, storage use and predictable limits.','The README reports a 19:02.88 sample export at 15.9 MB and about four minutes locally; that is one machine and one sample.'),
],[1.64,3.05,2.17],8.3)
src('Current metrics are from README.md checked 18 October 2026. Evaluation and feasibility advice: Workshop 1 pp. 11–13, 16, 21–22 and 24–25; Workshop 2 pp. 44–46.')
p('Do not claim that PDF overlays always use less compute or bandwidth than competitors. Measure actual export time, output size and cloud cost on comparable inputs. The workshop favors technical claims that are specific, real and supported by the submitted product.')
src('Workshop 1, pp. 13 and 16, 19:46–20:12 and 24:02–24:32; Workshop 2, pp. 44–46. Measurement proposal is project analysis.')

h('8  Proposal and demonstration package')
h('Proposal deck as PDF',2)
for x in [
    'Slide 1: learner, problem and chosen Media, Content & Digital Experiences theme.',
    'Slide 2: a concrete failure in plain audio with a long engineering PDF; show the missing visual context.',
    'Slide 3: one-screen product flow from PDF and narration to grounded timeline, page and mark.',
    'Slide 4: architecture showing local transcription/geometry, meaningful Google AI step, validation/review, Cloud Run or Firebase and storage.',
    'Slide 5: measured accuracy, coverage, false-mark count, latency, export size and known limitations; identify sample sizes.',
    'Slide 6: difference from existing overview or video tools, learner value, production scale and next validation step.',
]: b(x)
src('Workshop 2, pp. 43–44, timestamp not supplied; Workshop 1, pp. 20–22, 29:46–34:25. Slide sequence is this guide’s recommendation.')
h('Three-minute demo outline',2)
table(['Time','Show','Judging purpose'],[
('0:00–0:25','A student listens to long narration and cannot identify the referenced formula; open the working app immediately.','Problem and product clarity'),
('0:25–1:15','Use a prepared PDF and audio to show a correct page jump, text underline and equation/figure box.','Core value and usability'),
('1:15–1:55','Show one hard transcript passage, its source candidates and the Google AI contribution with validation or abstention.','Technical merit and meaningful Gen AI'),
('1:55–2:25','Show a marked uncertain case entering review and one correction; revisit a page with saved marks.','Reliability and differentiation'),
('2:25–3:00','Show a result metric, architecture snapshot, reachable Cloud URL and compact MP4 output.','Evidence, feasibility and close'),
],[.9,3.42,2.54],8.7)
p('The recorded product and live prototype must agree. Narration should be clear and audible. The Q&A speaker accepts edited or AI-assisted video production when it truthfully shows the working app, and explicitly warns that a polished video of nonexistent functionality works against a team. A live product opening is that speaker’s personal preference, not a judging rule.')
src('Workshop 1, pp. 20–22 and 25–26, 29:46–34:25 and 38:20–39:22; official demo length, https://aibuildercup.com/themes.html .')
h('Evaluator handoff',2)
for x in [
    'Provide the deployed URL, public demo URL, repository URL and proposal PDF in the submission form.',
    'Give one sample lesson and exact steps to reproduce three reliable grounded moments, plus any account details or usage limits.',
    'Use synthetic or licensed teaching material where appropriate and label synthetic data clearly.',
    'Test all links, playback, uploads and model calls from a clean browser; leave the app available throughout evaluation.',
    'Keep the submitted app stable after the deadline. The Q&A permits urgent fixes but advises against changes that disrupt judge access.',
]: b(x)
src('Workshop 1, pp. 3, 11, 19, 23–24 and 27–28, 04:32–05:31, 17:01–17:42, 29:05–29:30, 34:30–36:23 and 40:46–42:49.')

h('9  Risks and discrepancies to resolve')
table(['Item','What the sources say','Decision for the team'],[
('Student eligibility','Both workshops exclude students; current site says eligibility may be checked at any point.','Request a written determination from Hack2skill. Do not treat registration as approval of student status.'),
('Team dates and size','Intro video says up to two then up to four, and 11 October team deadline; later Q&A describes extension to 16 October. Current site says 2–4 by 14 October.','Check dashboard and organizer; do not substitute transcript dates for current operational instructions.'),
('Demo length','First Q&A appears to say three to four minutes; its presenter later defers to the published limit. Current site says three minutes.','Produce a three-minute demo unless dashboard gives a different explicit limit.'),
('Category wording','Official challenge page lists six themes but one submission bullet gives other example categories such as Education.','Choose Media, Content & Digital Experiences; confirm the exact dropdown on the submission form.'),
('Deployment interpretation','Q&A discusses mixed infrastructure, but official page requires a working Google Cloud deployment.','Make the submitted entry point and core user flow visibly run on Cloud Run or Firebase; document dependencies.'),
('Repository access','Q&A allows a minimal public evidence repo for proprietary work; other workshop text asks for a public repository.','Publish enough implementation evidence for judges, with no secrets; verify dashboard instructions.'),
],[1.36,3.18,2.32],8.3)
src('Workshop 1 pp. 2–4, 6–7 and 35; Workshop 2 pp. 36–37 and 43–45; official overview and themes pages checked 18 October 2026.')
p('Organizer contact listed on the current site: support+aibuildercup@hack2skill.com. Ask specifically whether currently enrolled engineering students may participate in this edition, even when registered or also employed. Preserve the written response.')

h('10  Source traceability')
p('W1 is the first Ask the Expert workshop (PDF pp. 1–35); W2 is the introductory workshop (beginning mid-p. 35). W1 transcript includes video timestamps. The supplied W2 transcript does not include usable video timestamps; the PDF page is therefore the precise available locator. The table covers each substantive topic, including lower-priority Q&A guidance.')
trace=[
('Eligibility and team membership','W1 p4 05:57–06:25; W2 p37','Official overview'),
('Team formation and submission dates','W1 pp2,35 ~03:00; 52:30–52:55; W2 p36','Official overview'),
('Proposal, demo, link, repository','W1 p3 04:32–05:31; W2 pp43–45','Official challenges'),
('Stay live through evaluation','W1 pp3,7–8,27 04:56–05:19; 11:05–12:40; 40:46–41:33','Official overview evaluation dates'),
('Cloud Run or Firebase deployment','W1 pp5,7,16–17,29–30 08:05–08:24; 10:16–10:59; 24:42–26:13; 43:44–46:39; W2 pp43–44','Official challenges'),
('Proprietary code and evidence repo','W1 pp6–7 08:41–10:08; W2 p45','Official challenges'),
('Cloud scaling and costs','W1 pp8,15–16 11:29–12:40; 22:47–24:32; W2 p47','—'),
('ADK versus SDK choice','W1 pp9–10 13:02–15:01; W2 pp47,49–50','—'),
('Grounding guardrails and evaluator data','W1 pp10–11 15:09–17:42','—'),
('Latency and hybrid architecture','W1 pp12–14 18:01–22:07','Official technical merit'),
('Language, timeout, secrets and quotas','W1 pp14–19 22:07–29:30','—'),
('Truthful demo and audible explanation','W1 pp19–22 29:46–34:25','Official demo requirement'),
('Login friction and synthetic examples','W1 pp23–24 34:30–36:23','—'),
('Prototype feasibility and judge viewpoint','W1 pp24–26 36:23–39:22','Official technical merit'),
('503 contingency and stable test examples','W1 pp26–28 39:22–42:49','—'),
('Split infrastructure dependencies','W1 pp29–31 43:44–47:08','Official deployment rule'),
('Local model plus Gemini; database choice','W1 pp31–34 47:13–51:28; W2 pp48–50','Official AI requirement'),
('Six themes and education fit','W2 pp38–42','Official themes'),
('Functional prototype and documentation','W2 pp42–45','Official challenges'),
('Judging criteria and innovation','W2 pp45–46','Official rubric'),
('Simple architecture and RAG choices','W2 pp47–50','—'),
]
table(['Topic','Workshop location','Rule cross-check'],trace,[2.25,3.36,1.25],8.1)
p('Official pages checked on 18 October 2026: https://aibuildercup.com/ and https://aibuildercup.com/themes.html . Project-state source: C:\\Users\\User\\Desktop\\Vscode\\Audio-to-PDF\\README.md and the current backend modules. Source transcript: C:\\Users\\User\\Downloads\\Untitled document (6).pdf .')

foot=sec.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=foot.add_run('Audio to PDF  |  Workshop analysis  |  18 October 2026  |  '); run.font.size=Pt(8); run.font.color.rgb=RGBColor(90,99,110)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); foot._p.append(fld)

doc.save(OUT)
print(OUT)
